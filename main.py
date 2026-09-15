from kickbase_api.kickbase import Kickbase
import os
import requests
import pandas as pd
import openpyxl

kickbase = Kickbase()

# gets user credentials from environment variables and logs in
USERNAME = os.environ["kickbase_user"]
PASSWORD = os.environ["kickbase_pw"]
url = "https://api.kickbase.com/v4/user/login"
headers = {
    "Content-Type": "application/json",
    "Accept": "application/json"
}
payload = {
    "em": USERNAME,
    "pass": PASSWORD,
    "ext": True,  # TODO: What is this?
    "loy": False,  # TODO: What is this?
    "rep": {}  # TODO: What is this?
}
# token beschaffen für weiteren Zugriff
response = requests.post(url, json=payload, headers=headers).json()
token = response["tkn"]

# ID der Liga beschaffen
leagues_endpoint = "https://api.kickbase.com/v4/leagues/selection"
leagues_headers = {
        "Content-Type": "application/json",
        "Accept": "application/json",
        "Authorization": f"Bearer {token}"
    }

leagues = requests.request("GET", leagues_endpoint, headers=leagues_headers, data=payload)
leagueID = leagues.json()['it'][2]['i']

# Spielerdaten ziehen
players_endpoint = f"https://api.kickbase.com/v4/leagues/{leagueID}/squad"
players_headers = {
        "Content-Type": "application/json",
        "Accept": "application/json",
        "Authorization": f"Bearer {token}"
}

players_response = requests.request("GET", players_endpoint, headers=players_headers, data=payload)
players_data = players_response.json()
print(players_data)
players_list = [(player['n'], player['mv'], player.get('ap', 0), player['pos']) for player in players_data['it']]
print(players_list)

playersID = [player['i'] for player in players_data['it']]
print(playersID)
# einkaufspreis = []
# for ID in playersID:
#     transferHistory_url = f"https://api.kickbase.com/v4/leagues/{leagueID}/players/{ID}/transferHistory"
#     response_ek = requests.request("GET", transferHistory_url, headers=leagues_headers, data=payload)
#     ek = response_ek.json()['it'][-1]['trp']
#     einkaufspreis.append(ek)

# # Budgetdaten ziehen

budget_endpoint = f"https://api.kickbase.com/v4/leagues/{leagueID}/me/budget"
budget_headers = {
        "Content-Type": "application/json",
        "Accept": "application/json",
        "Authorization": f"Bearer {token}"
}

budget_response = requests.request("GET", budget_endpoint, headers=budget_headers, data=payload)
budget_data = budget_response.json()
current_cash = budget_data['b']
print(budget_data)

# # Spielerdaten nach Excel exportieren
df = pd.DataFrame(players_list, columns=['Spielername', 'Marktwert', 'Punkte Durchschnitt', 'Position'])
# df["Einkaufspreis"] = einkaufspreis
# df["Differenz"] = df["Marktwert"] - df["Einkaufspreis"]
#

df = df.sort_values(by=["Position", "Spielername"])
positions_map = {
    1: "Torwart",
    2: "Abwehr",
    3: "Mittelfeld",
    4: "Sturm"
}
df["Position"] = df["Position"].map(positions_map)
with pd.ExcelWriter("kickbase_spieler.xlsx", engine="openpyxl") as writer:
    df.to_excel(writer, index=False, sheet_name="Spieler")

    workbook = writer.book
    worksheet = writer.sheets["Spieler"]
    # Zahlenformat setzen
    for cell in worksheet["B"][1:]:  # [1:] überspringt Header
        cell.number_format = "#,##0"

    # Zwei Leerzeilen und anschließend Kontostand
    start_row = len(df) + 4
    worksheet.cell(row=start_row, column=1).value = "Kontostand:"
    worksheet.cell(row=start_row, column=2).value = current_cash
    worksheet.cell(row=start_row, column=2).number_format = "#,##0"

